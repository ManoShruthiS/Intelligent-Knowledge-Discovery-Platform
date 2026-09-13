import io
import csv
import json
import logging
import os
import zipfile
import fitz
from docx import Document
from services.pii_redactor_service import pii_redactor_service
logger = logging.getLogger(__name__)
PII_REDACT_ENABLED = os.getenv('PII_REDACT', '0') == '1'

class UnsupportedFileTypeError(Exception):
    pass

class DocumentProcessingError(Exception):
    pass
CODE_EXTENSIONS = {'.py': 'Python', '.js': 'JavaScript', '.jsx': 'JavaScript', '.ts': 'TypeScript', '.tsx': 'TypeScript', '.java': 'Java', '.c': 'C', '.h': 'C', '.cpp': 'C++', '.hpp': 'C++', '.cc': 'C++', '.go': 'Go', '.rs': 'Rust', '.rb': 'Ruby', '.sh': 'Shell', '.bash': 'Shell', '.html': 'HTML', '.htm': 'HTML', '.css': 'CSS', '.xml': 'XML', '.yaml': 'YAML', '.yml': 'YAML', '.toml': 'TOML'}
TEXT_EXTENSIONS = {'.txt', '.md', '.markdown'}
DATA_EXTENSIONS = {'.csv', '.json'}
IMAGE_EXTENSIONS = {'.png', '.jpg', '.jpeg', '.gif', '.webp', '.bmp'}
NOTEBOOK_EXTENSIONS = {'.ipynb'}

def process_document(file_content: bytes, filename: str) -> dict:
    if not file_content:
        raise DocumentProcessingError('File is empty.')
    ext_dot = _ext(filename)
    ext = ext_dot.lstrip('.') if ext_dot else ''
    file_type = ''
    pages = []
    extra_meta = {}
    author = None
    year = None
    try:
        if ext == 'pdf':
            file_type = 'PDF'
            pages = _extract_pdf(file_content)
        elif ext == 'docx':
            file_type = 'DOCX'
            pages = _extract_docx(file_content)
        elif ext_dot in TEXT_EXTENSIONS:
            file_type = 'Markdown' if ext.startswith('m') else 'TXT'
            pages = _extract_text(file_content)
        elif ext == 'csv':
            file_type = 'CSV'
            pages = _extract_csv(file_content)
        elif ext == 'xlsx':
            file_type = 'XLSX'
            pages, extra_meta = _extract_xlsx(file_content)
        elif ext == 'json':
            file_type = 'JSON'
            pages = _extract_json(file_content)
        elif ext_dot in CODE_EXTENSIONS:
            file_type = CODE_EXTENSIONS[ext_dot]
            pages = _extract_text(file_content)
        elif ext == 'ipynb':
            file_type = 'Jupyter Notebook'
            pages = _extract_ipynb(file_content)
        elif ext_dot in IMAGE_EXTENSIONS:
            file_type = ext.upper().lstrip('.')
            pages = _extract_image(file_content, filename)
        else:
            raise UnsupportedFileTypeError(f"Unsupported file type: {ext_dot or '(none)'}")
    except UnsupportedFileTypeError:
        raise
    except Exception as e:
        logger.error(f'Extraction failed for {filename}: {e}')
        raise DocumentProcessingError(f'Failed to extract text from {filename}: {e}')
    full_text = '\n\n'.join((p['text'] for p in pages if p.get('text')))
    if file_type == 'PDF':
        try:
            import fitz
            doc = fitz.open(stream=file_content, filetype='pdf')
            page_count = len(doc)
            doc.close()
        except Exception:
            page_count = len(pages)
    else:
        page_count = None
    pii_redactions = 0
    if PII_REDACT_ENABLED and full_text:
        redacted_text = pii_redactor_service.redact_text(full_text)
        pii_redactions = sum((redacted_text.count(token) for token in ('[REDACTED_EMAIL]', '[REDACTED_PHONE]', '[REDACTED_CC]', '[REDACTED_SSN]', '[REDACTED_IP]')))
        if pii_redactions:
            logger.info(f'PII redactor stripped {pii_redactions} item(s) from {filename}')
            full_text = redacted_text
            redacted_pages = []
            for page in pages:
                page_text = page.get('text') or ''
                if not page_text.strip():
                    redacted_pages.append(page)
                    continue
                replaced = pii_redactor_service.redact_text(page_text)
                redacted_pages.append({**page, 'text': replaced})
            pages = redacted_pages
    return {'filename': filename, 'file_type': file_type, 'file_size': len(file_content), 'page_count': page_count, 'character_count': len(full_text), 'text': full_text, 'pages': pages, 'metadata': extra_meta, 'author': author, 'year': year, 'pii_redacted': pii_redactions if PII_REDACT_ENABLED else 0}

def _extract_pdf(file_content: bytes) -> list[dict]:
    import tempfile
    import gc
    import os
    import sys
    kno_deps_path = 'C:\\Users\\manos\\kno_deps'
    if kno_deps_path not in sys.path:
        sys.path.insert(0, kno_deps_path)
    try:
        from docling.document_converter import DocumentConverter
    except ImportError:
        logger.warning('Docling is not installed. Falling back to PyMuPDF.')
        return _extract_pdf_pymupdf(file_content)
    with tempfile.NamedTemporaryFile(suffix='.pdf', delete=False) as tf:
        tf.write(file_content)
        temp_path = tf.name
    try:
        converter = DocumentConverter()
        result = converter.convert(temp_path)
        md_text = result.document.export_to_markdown()
        return [{'page_number': None, 'text': md_text}]
    except Exception as e:
        logger.error(f'Docling conversion failed: {e}')
        return _extract_pdf_pymupdf(file_content)
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)
        gc.collect()

def _extract_pdf_pymupdf(file_content: bytes) -> list[dict]:
    pages = []
    doc = fitz.open(stream=file_content, filetype='pdf')
    try:
        for i in range(len(doc)):
            page = doc.load_page(i)
            text = page.get_text() or ''
            spans: list[dict] = []
            if text.strip():
                cursor = 0
                try:
                    page_dict = page.get_text('dict')
                    for block in page_dict.get('blocks', []):
                        if block.get('type') != 0:
                            continue
                        for line in block.get('lines', []):
                            for span in line.get('spans', []):
                                span_text = span.get('text', '')
                                if not span_text:
                                    continue
                                bbox = span.get('bbox')
                                if not bbox:
                                    continue
                                spans.append({'char_start': cursor, 'char_end': cursor + len(span_text), 'bbox': [float(bbox[0]), float(bbox[1]), float(bbox[2]), float(bbox[3])]})
                                cursor += len(span_text)
                except Exception:
                    spans = []
                pages.append({'page_number': i + 1, 'text': text, 'spans': spans, 'width': float(page.rect.width), 'height': float(page.rect.height)})
    finally:
        doc.close()
    return pages

def _extract_docx(file_content: bytes) -> list[dict]:
    doc = Document(io.BytesIO(file_content))
    paragraphs = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    tables_text = []
    for t in doc.tables:
        rows = []
        for row in t.rows:
            rows.append(' | '.join((cell.text.strip() for cell in row.cells)))
        if rows:
            tables_text.append('\n'.join(rows))
    body = '\n'.join(paragraphs)
    if tables_text:
        body += '\n\n' + '\n\n'.join(tables_text)
    return [{'page_number': None, 'text': body}]

def _extract_text(file_content: bytes) -> list[dict]:
    for enc in ('utf-8', 'latin-1', 'windows-1252'):
        try:
            return [{'page_number': None, 'text': file_content.decode(enc)}]
        except UnicodeDecodeError:
            continue
    raise DocumentProcessingError('Unable to decode text file.')

def _extract_csv(file_content: bytes) -> list[dict]:
    text = _decode_bytes(file_content)
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    if not rows:
        return [{'page_number': None, 'text': ''}]
    header = rows[0]
    lines = [', '.join(header)]
    for r in rows[1:]:
        lines.append(', '.join(r))
    return [{'page_number': None, 'text': '\n'.join(lines)}]

def _extract_xlsx(file_content: bytes) -> tuple[list[dict], dict]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise DocumentProcessingError('XLSX support requires openpyxl. Run: pip install openpyxl')
    wb = load_workbook(io.BytesIO(file_content), read_only=True, data_only=True)
    pages = []
    meta = {'sheet_names': []}
    for sheet_name in wb.sheetnames:
        meta['sheet_names'].append(sheet_name)
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = ['' if v is None else str(v) for v in row]
            if any((c.strip() for c in cells)):
                rows.append(' | '.join(cells))
        if rows:
            pages.append({'page_number': None, 'text': f'# Sheet: {sheet_name}\n' + '\n'.join(rows)})
    wb.close()
    return (pages, meta)

def _extract_json(file_content: bytes) -> list[dict]:
    text = _decode_bytes(file_content)
    try:
        data = json.loads(text)
        flat = json.dumps(data, indent=2, ensure_ascii=False)
    except Exception:
        flat = text
    return [{'page_number': None, 'text': flat}]

def _extract_ipynb(file_content: bytes) -> list[dict]:
    text = _decode_bytes(file_content)
    try:
        nb = json.loads(text)
    except Exception:
        return [{'page_number': None, 'text': text}]
    out = []
    cells = nb.get('cells', [])
    for i, c in enumerate(cells):
        ctype = c.get('cell_type', '')
        src = ''.join(c.get('source', [])) if isinstance(c.get('source'), list) else str(c.get('source', ''))
        if ctype == 'markdown':
            out.append({'page_number': None, 'text': f'[Cell {i} — Markdown]\n{src}'})
        elif ctype == 'code':
            out.append({'page_number': None, 'text': f'[Cell {i} — Code]\n{src}'})
        else:
            out.append({'page_number': None, 'text': f'[Cell {i} — {ctype}]\n{src}'})
    if not out:
        out = [{'page_number': None, 'text': json.dumps(nb, indent=2)}]
    return out

def _extract_image(file_content: bytes, filename: str) -> list[dict]:
    note = f'[Image attachment: {filename}]\nThe text content of this file is not extractable. When used as an attachment in a chat, the model with vision capability may analyse its visual content.'
    return [{'page_number': None, 'text': note}]

def _ext(filename: str) -> str:
    if '.' not in filename:
        return ''
    return '.' + filename.rsplit('.', 1)[-1].lower()

def _decode_bytes(b: bytes) -> str:
    for enc in ('utf-8', 'latin-1', 'windows-1252'):
        try:
            return b.decode(enc)
        except UnicodeDecodeError:
            continue
    raise DocumentProcessingError('Unable to decode bytes.')
